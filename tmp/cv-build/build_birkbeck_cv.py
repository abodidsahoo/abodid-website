from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = ROOT / "artifacts" / "birkbeck-application"
OUTPUT_PATH = OUTPUT_DIR / "Abodid-Sahoo-Birkbeck-Senior-Technologist-CV.docx"

BLACK = RGBColor(0x00, 0x00, 0x00)
NAVY = RGBColor(0x08, 0x1F, 0x5C)
MUTED = RGBColor(0x4B, 0x55, 0x63)
LIGHT = "D9DEE8"


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=LIGHT, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_run_font(run, name="Arial"):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)


def add_hyperlink(paragraph, text, url, color="244BCE", underline=True):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "19")
    r_pr.append(sz)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.25)
    normal.font.color.rgb = BLACK
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.paragraph_format.space_after = Pt(4.2)
    normal.paragraph_format.line_spacing = 1.06

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.keep_with_next = True
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 13.2, 10, 4),
        ("Heading 2", 11.1, 5, 1.5),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "CV Subtitle" not in styles:
        subtitle = styles.add_style("CV Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        subtitle = styles["CV Subtitle"]
    subtitle.font.name = "Arial"
    subtitle.font.size = Pt(13.4)
    subtitle.font.bold = False
    subtitle.font.color.rgb = NAVY
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    subtitle.paragraph_format.space_after = Pt(4)
    subtitle.paragraph_format.keep_with_next = True

    if "CV Meta" not in styles:
        meta = styles.add_style("CV Meta", WD_STYLE_TYPE.PARAGRAPH)
    else:
        meta = styles["CV Meta"]
    meta.font.name = "Arial"
    meta.font.size = Pt(9.4)
    meta.font.color.rgb = MUTED
    meta._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    meta._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    meta.paragraph_format.space_after = Pt(7)

    if "CV Bullet" not in styles:
        bullet = styles.add_style("CV Bullet", WD_STYLE_TYPE.PARAGRAPH)
    else:
        bullet = styles["CV Bullet"]
    bullet.base_style = styles["Normal"]
    bullet.font.name = "Arial"
    bullet.font.size = Pt(9.85)
    bullet._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    bullet._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    bullet.paragraph_format.left_indent = Inches(0.18)
    bullet.paragraph_format.first_line_indent = Inches(-0.14)
    bullet.paragraph_format.space_after = Pt(2.6)
    bullet.paragraph_format.line_spacing = 1.03


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run(text.upper())
    set_run_font(r)
    return p


def add_role_heading(doc, role, organization, dates):
    p = doc.add_paragraph(style="Heading 2")
    r = p.add_run(role)
    r.bold = True
    set_run_font(r)
    r2 = p.add_run(f"  |  {organization}")
    r2.bold = False
    r2.font.color.rgb = MUTED
    set_run_font(r2)
    r3 = p.add_run(f"  |  {dates}")
    r3.bold = False
    r3.font.color.rgb = MUTED
    set_run_font(r3)
    return p


def add_bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="CV Bullet")
    mark = p.add_run("•  ")
    mark.bold = True
    mark.font.color.rgb = NAVY
    set_run_font(mark)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        lead.bold = True
        set_run_font(lead)
        body = p.add_run(text[len(bold_lead):])
        set_run_font(body)
    else:
        body = p.add_run(text)
        set_run_font(body)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.22)
    configure_styles(doc)

    # Footer is deliberately minimal and repeated for navigation only.
    footer_p = section.footer.paragraphs[0]
    footer_p.text = ""
    add_page_number(footer_p)

    title = doc.add_paragraph(style="Title")
    title.add_run("Abodid Sahoo")
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    subtitle = doc.add_paragraph(style="CV Subtitle")
    subtitle.add_run("Creative Technologist  |  AI Systems  |  Digital Media and Immersive Production")
    meta = doc.add_paragraph(style="CV Meta")
    meta.add_run("London and Bengaluru  |  ")
    add_hyperlink(meta, "abodid@network.rca.ac.uk", "mailto:abodid@network.rca.ac.uk")
    meta.add_run("  |  +44 752 225 8768  |  ")
    add_hyperlink(meta, "abodid.com", "https://www.abodid.com/")
    meta.add_run("  |  ")
    add_hyperlink(meta, "LinkedIn", "https://www.linkedin.com/in/abodidsahoo/")

    add_section_heading(doc, "Profile")
    p = doc.add_paragraph()
    p.add_run(
        "Creative technologist, filmmaker and researcher with 9+ years of experience across "
        "AI-enabled systems, film, photography, exhibitions and public-facing media. I build "
        "working creative technology rather than concept-only demonstrations: retrieval and "
        "semantic-search systems, multimodal vision experiments, browser-based computer vision, "
        "automated research pipelines and accessible interactive tools."
    )
    p2 = doc.add_paragraph()
    p2.add_run(
        "At the Royal College of Art, I led media delivery for graduate exhibitions across more "
        "than 10 spaces and 700 international students, working across AV setup, lighting, digital "
        "media installations, XR, spatial layout and public presentation. I also teach creative "
        "workflows to non-specialist audiences and document technical systems so that other people "
        "can use and maintain them."
    )

    add_section_heading(doc, "Relevant technical capabilities")
    capabilities = [
        ("Creative AI systems", "OpenRouter, Gemini, Claude and OpenAI APIs; multimodal image analysis and generation; RAG; embeddings; pgvector; structured outputs; model routing; retries and deterministic validation."),
        ("Software and infrastructure", "TypeScript, JavaScript, React, Astro, Python data-processing scripts, command-line workflows, Postgres, Supabase, GitHub, Vercel, Cloudflare R2, REST APIs and automated jobs."),
        ("Computer vision and interaction", "MediaPipe, TensorFlow.js, WebGL and Canvas; real-time hand tracking; accessible keyboard, touch and mouse fallbacks; spatial and XR prototyping."),
        ("Responsible operation", "API-key separation, rate limits, usage and cost controls, privacy-by-default ingestion, data provenance, public/private access rules, citation validation, accessibility and health-and-safety practice."),
        ("Media and teaching", "Video production, cinematography, editing, photography, AV setup, lighting, exhibition installation, learner support, workshop delivery and technical documentation."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.68)
    table.columns[1].width = Inches(5.55)
    hdr = table.rows[0].cells
    hdr[0].text = "Area"
    hdr[1].text = "Evidence and tools"
    for cell in hdr:
        set_cell_fill(cell, "E8ECF5")
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = BLACK
                run.font.size = Pt(9.4)
                set_run_font(run)
    set_repeat_table_header(table.rows[0])
    for idx, (label, detail) in enumerate(capabilities):
        row = table.add_row().cells
        row[0].text = label
        row[1].text = detail
        if idx % 2:
            for cell in row:
                set_cell_fill(cell, "F6F8FB")
        for cell in row:
            set_cell_margins(cell)
            cell.vertical_alignment = 1
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(9.1)
                    set_run_font(run)
        row[0].paragraphs[0].runs[0].bold = True
    set_table_borders(table)

    add_section_heading(doc, "Selected creative AI and interactive systems")
    add_role_heading(doc, "Obsidian Knowledge Vault", "RAG and semantic search", "2025-2026")
    add_bullet(doc, "Designed a local-to-public ingestion pipeline for an Obsidian archive, with explicit path allowlists, sensitive-content detection, semantic chunking, SHA-256 change detection and batched embedding generation.")
    add_bullet(doc, "Implemented 1,536-dimension pgvector storage, hybrid vector and keyword retrieval, multi-model routing and cited conversational answers. Added deterministic citation remapping after finding that unconstrained model output produced invalid source links.")
    add_bullet(doc, "Operated commercial model APIs with server-side keys, rate limiting, lower-cost model selection, batch processing and documented service quotas to control spend and downtime.")

    add_role_heading(doc, "Punctum Research Suite", "Multimodal visual research", "2025-2026")
    add_bullet(doc, "Built a participatory study comparing human attention in photographs with multimodal model interpretations. The system records normalized coordinates, emotion labels and participant reasoning, then compares human clusters with model-generated attention regions.")
    add_bullet(doc, "Implemented anonymous sessions, Cloudflare Turnstile validation, generation limits and model fallbacks. The research treats bias, privacy, consent and the limits of machine interpretation as part of the method, not an afterthought.")

    # Force a stable second page so the CV remains easy to edit in Word.
    doc.add_page_break()

    add_section_heading(doc, "Selected systems continued")
    add_role_heading(doc, "Image Flick and Hand Gesture Control", "Browser computer vision", "2025-2026")
    add_bullet(doc, "Built a touch-free photograph browser using MediaPipe, TensorFlow.js and a WebGL backend to track hand position and gesture velocity in real time. Added touch and mouse fallbacks for unsupported hardware and different access needs.")

    add_role_heading(doc, "Sequence Room", "Spatial narrative authoring", "2026")
    add_bullet(doc, "Created an interactive canvas for arranging, rotating, grouping and annotating photographs as a visual sequence. Built with React, dnd-kit and HTML Canvas, with keyboard-accessible controls and PDF storyboard export.")

    add_role_heading(doc, "Research Automation", "Model evaluation and operational pipelines", "2026")
    add_bullet(doc, "Built a daily reading pipeline that discovers, verifies, deduplicates and ranks sources before sending exactly five items, with idempotent delivery and audit records. Built an Opportunity Radar that combines institutional web scraping with structured model extraction and deadline reminders.")

    add_section_heading(doc, "Professional experience")
    add_role_heading(doc, "Exhibition and Media Technical Lead", "Royal College of Art, London", "2023-2025")
    add_bullet(doc, "Led media delivery across 10+ graduate exhibition spaces supporting 700+ international students. Coordinated AV setup, lighting, digital installations, spatial layouts, technical troubleshooting and public presentation while maintaining health-and-safety requirements.")
    add_bullet(doc, "Supported exhibition teams and students with different levels of technical confidence, translating creative intentions into workable media and installation plans under live delivery deadlines.")

    add_role_heading(doc, "Creative Producer, Director and Editor", "Independent and client work", "2018-2025")
    add_bullet(doc, "Delivered films, photography and visual storytelling for organisations including the British Film Institute, Outernet London, Frameless, Wieden+Kennedy, Budweiser, Uniqlo, Odisha Tourism, Hermosa Design Studio and Pursuit of Portraits in New York.")
    add_bullet(doc, "Worked end to end across research, treatments, budgets, production planning, direction, cinematography, editing, visual effects and final delivery. Work was broadcast on VH1 India and published by Rolling Stone India, The Indian Express and Homegrown.")

    add_role_heading(doc, "Founder and Creative Director", "Visual Notes, India", "2014-2017")
    add_bullet(doc, "Led a small creative practice delivering UX design, photography, promotional media, brand films and identity work, coordinating clients, collaborators and production schedules.")

    add_section_heading(doc, "Teaching research and leadership")
    add_bullet(doc, "Delhi University: guest lecture and practical workshop on ethnographic filmmaking for 50+ BA students.")
    add_bullet(doc, "British Sociological Association 2026: presented two research projects and chaired two conference sessions with 100+ international delegates.")
    add_bullet(doc, "Cambridge Digital Humanities Cultural Heritage Data School 2026: full bursary recipient; worked with cultural data, Python, photogrammetry, critical visualisation and accessibility, and presented to 50+ fellows from 15+ countries.")

    add_section_heading(doc, "Education and recognition")
    entries = [
        ("Royal College of Art, London", "MA Digital Direction, 2022-2023  |  Apple Scholarship, full tuition award of £28,000"),
        ("National Institute of Design, Ahmedabad", "MDes Film and Video Communication, 2021-2022  |  All India Rank 1; moved to RCA after Year 1"),
        ("Prahlad Kakar School of Branding and Entrepreneurship", "Diploma in Filmmaking, Branding and Entrepreneurship, 2018-2019"),
        ("National Institute of Technology, Rourkela", "BTech and MTech Mechanical Engineering, 2012-2017"),
    ]
    for institution, detail in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2.2)
        a = p.add_run(institution + "  |  ")
        a.bold = True
        set_run_font(a)
        b = p.add_run(detail)
        set_run_font(b)

    add_section_heading(doc, "Portfolio")
    p = doc.add_paragraph(style="CV Meta")
    add_hyperlink(p, "Selected work", "https://www.abodid.com/")
    p.add_run("  |  ")
    add_hyperlink(p, "Research", "https://www.abodid.com/research")
    p.add_run("  |  ")
    add_hyperlink(p, "Architecture", "https://www.abodid.com/architecture")
    p.add_run("  |  ")
    add_hyperlink(p, "Films", "https://www.abodid.com/films")

    # Accessibility metadata and document properties.
    doc.core_properties.title = "Abodid Sahoo Senior Technologist CV"
    doc.core_properties.subject = "Application for Senior Technologist Birkbeck Centre for Creative AI"
    doc.core_properties.author = "Abodid Sahoo"
    doc.core_properties.keywords = "creative AI, senior technologist, digital education, media production, immersive technology"

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
